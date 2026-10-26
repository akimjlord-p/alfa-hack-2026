"""Build editable DOCX documentation directly from the current questionnaire code."""
from __future__ import annotations

from datetime import date
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.questionnaire import QUESTIONS, PURPOSE_CHOICES_BY_INDUSTRY, VERSION, choices_for_question  # noqa: E402

OUT = ROOT / "docs"
GREEN = RGBColor(13, 96, 69)
DARK = RGBColor(21, 38, 30)
MUTED = RGBColor(90, 111, 99)


def base_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = DARK
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = GREEN
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.color.rgb = GREEN
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 2"].font.color.rgb = GREEN
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.style.font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("akimai  •  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def paragraph(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(value)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = head
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "0D6045")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


FLOW = [
    ("industry", "Всем"),
    ("purpose", "Всем"),
    ("agri_calendar", "Только если отрасль — сельское хозяйство"),
    ("marketplace_cycle", "Только если отрасль — маркетплейсы"),
    ("contract_terms", "Если цель — коммерческий контракт или госконтракт"),
    ("asset_status", "Если цель — оборудование, транспорт, коммерческая недвижимость, инвестиции или строительство"),
    ("urgency", "Если цель — кассовый разрыв"),
    ("gap_source", "Сразу после вопроса о срочности кассового разрыва"),
    ("spending_pattern", "Всем"),
    ("first_tranche", "Если расходы идут 2–3 этапами или ежемесячно"),
    ("supplier_deferral", "Если цель — сезонные или оборотные закупки, оборудование, транспорт или маркетплейс"),
    ("amount", "Всем"),
    ("inflow", "Всем"),
    ("business_age", "Всем"),
    ("legal_form", "Всем"),
    ("revenue", "Всем"),
    ("cash_surplus", "Всем"),
    ("collateral", "Всем"),
    ("debt", "Всем"),
]


def build_questionnaire() -> Path:
    doc = base_doc(
        "Опросник akimai — точный маршрут",
        f"Реализованная схема · версия {VERSION} · {date.today():%d.%m.%Y}",
    )
    doc.add_paragraph(
        "Вопросы ниже приведены дословно и в порядке, в котором их проверяет сервер. "
        "Условный вопрос задаётся только при выполнении его условия. После каждого ответа маршрут пересчитывается."
    )
    doc.add_heading("Последовательность и ветвления", level=1)
    for number, (qid, condition) in enumerate(FLOW, 1):
        q = QUESTIONS[qid]
        doc.add_heading(f"{number}. {q.title}", level=2)
        paragraph(doc, "Когда: ", condition + ".")
        paragraph(doc, "Быстрые ответы: ", " · ".join(choice.label for choice in q.choices) + ".")
        if any(choice.input_required for choice in q.choices):
            names = ", ".join(f"«{choice.label}»" for choice in q.choices if choice.input_required)
            paragraph(doc, "Свободный ввод: ", f"{names} открывает поле ответа и не завершает вопрос до отправки текста.")
        if qid == "purpose":
            doc.add_paragraph("После ответа об отрасли кнопки цели показываются в таком порядке:")
            table(doc, ["Отрасль", "Быстрые ответы на вопрос о цели"], [
                [industry, " · ".join(choice.label for choice in choices_for_question(q, {"industry": industry}))]
                for industry in PURPOSE_CHOICES_BY_INDUSTRY
            ])
    doc.add_heading("Когда вопрос считается отвеченным", level=1)
    doc.add_paragraph(
        "Обычная кнопка записывает заранее заданные факты. Свободный текст передаётся извлечению фактов: "
        "модель Yandex AI Studio получает текущий вопрос, уже известный профиль и последнее сообщение; каждый факт "
        "должен подтверждаться фрагментом именно этого сообщения. Вызов идёт через "
        "https://ai.api.cloud.yandex.net/v1 с моделью gpt://<folder_id>/yandexgpt/latest. Сервер проверяет "
        "допустимые поля и диапазоны. После ответа модели запускается локальное безопасное распознавание очевидных "
        "фраз: например, «нашел объект» означает asset_status=ready, «пока ищем объект» — asset_status=search, "
        "«оплата по этапам» — asset_status=staged. Локальный результат только дополняет пустые поля модели. "
        "Если провайдер недоступен, JSON невалиден или запрос длится более 8 секунд, используется этот fallback. "
        "Если нужное поле не определено, текущий вопрос остаётся открытым. Для нестандартного агрографика текст "
        "сохраняется в cashflow_details, а cashflow_cycle становится unknown, если ИИ не смог его классифицировать."
    )
    doc.add_heading("Завершение и пример", level=1)
    doc.add_paragraph(
        "Подбор запускается после завершения вопросов текущей ветки и получения суммы, цели, срока до возврата "
        "выручки, возраста бизнеса, юридической формы, выручки и обеспечения. Некоторые вопросы могут быть пропущены, "
        "если соответствующий факт уже достоверно извлечён из более раннего свободного ответа."
    )
    doc.add_paragraph(
        "Пример: фермер → «Семена, удобрения, сезонные работы» → «Нет, 2–3 закупки». "
        "Порядок номеров: 1 → 2 → 3 → 9 → 10 → 11 → 12 → 13 → 14 → 15 → 16 → 17 → 18 → 19."
    )
    paragraph(doc, "Источник реализации: ", "app/questionnaire.py, app/main.py, app/ai.py.")
    path = OUT / "anketa_akimai.docx"
    doc.save(path)
    return path


def build_ranking() -> Path:
    doc = base_doc(
        "akimai — как сортируются кредиты",
        f"Фактический алгоритм в текущем коде · {date.today():%d.%m.%Y}",
    )
    doc.add_heading("1. Что попадает в расчёт", level=1)
    doc.add_paragraph(
        "Профиль анкеты содержит цель, сумму, обеспечение и сведения о бизнесе. Расчётный срок кредита = "
        "max(3, min(240, 2 × число месяцев до начала возврата выручки)). Это техническое допущение, "
        "а не выбранный пользователем срок. Каталог сейчас содержит 12 заранее внесённых предложений конкретных банков."
    )
    doc.add_heading("2. Обязательные фильтры", level=1)
    table(doc, ["Фильтр", "Правило"], [
        ["Цель", "Цель пользователя должна входить в список purposes продукта."],
        ["Сумма", "min_amount ≤ запрошенная сумма ≤ max_amount."],
        ["Обеспечение", "Выбранный тип обеспечения должен входить в список collateral продукта."],
    ])
    doc.add_paragraph(
        "Если хотя бы один фильтр не пройден, продукт исключается. Возраст бизнеса, расчётный срок, "
        "юридическая форма, выручка и просрочка сейчас не являются жёсткими фильтрами."
    )
    doc.add_heading("3. Определение подходящего типа инструмента", level=1)
    doc.add_paragraph(
        "Тип предложения определяется по названию: «овердрафт» → overdraft; «линия» → credit_line; "
        "«контракт»/«госзаказ» → contract_financing; «ипотека»/«залог приобретаемого»/«капитальные»/" 
        "«инвестиционный» → asset_financing; иначе term_loan. Проверки идут именно в этом порядке."
    )
    table(doc, ["Параметры анкеты", "Предпочтение по правилам"], [
        ["Транши, регулярные расходы или сезонная цель", "Кредитная линия"],
        ["Кассовый разрыв и возврат выручки ≤ 3 месяцев", "Овердрафт и кредитная линия"],
        ["Коммерческий/госконтракт при наличии контракта", "Контрактное финансирование"],
        ["Оборудование, транспорт, недвижимость, стройка, инвестиции", "Финансирование актива"],
        ["Если ничего из перечисленного не подошло", "Разовый кредит"],
    ])
    doc.add_heading("4. Формула баллов", level=1)
    doc.add_paragraph("Итог = инструмент + денежный цикл + условия каталога + платёжная нагрузка + скорость + контекст ИИ. Максимум — 100 баллов.")
    table(doc, ["Компонент", "Точный расчёт", "Максимум"], [
        ["Инструмент", "22 базовых + 13, если тип входит в предпочтения правил.", "35"],
        ["Денежный цикл", "12, если расчётный срок внутри вилки продукта; иначе 0. Ещё 8 за совпадение: короткий кассовый разрыв/овердрафт, подтверждённый контракт/контрактное финансирование, транши/линия либо выбранный или поэтапно оплачиваемый актив/финансирование актива.", "20"],
        ["Условия каталога", "16 базовых + 4, если возраст бизнеса ≥ минимального возраста продукта. Если меньше — предупреждение, но продукт остаётся в ленте.", "20"],
        ["Платёжная нагрузка", "При положительном свободном остатке: 15, если оценочный платёж ≤ 70% остатка; 10, если ≤ 100%; иначе 0. Без положительного остатка — 5 и предупреждение.", "15"],
        ["Скорость", "При известной срочности и сроке рассмотрения: 5, если успевает, иначе 0. Если одного из параметров нет — 3.", "5"],
        ["Контекст ИИ", "5, только если ИИ предложил этот тип и он уже входит в предпочтения правил. Иначе 0.", "5"],
    ])
    doc.add_paragraph(
        "Оценочный платёж — аннуитет по запрошенной сумме и расчётному сроку при фиксированных 35% годовых. "
        "Это стресс-тест, а не ставка, ПСК или платёж по конкретному продукту."
    )
    doc.add_heading("5. Роль ИИ и порядок карточек", level=1)
    doc.add_paragraph(
        "Yandex AI Studio получает ограниченный набор фактов профиля через OpenAI-совместимый endpoint "
        "https://ai.api.cloud.yandex.net/v1 и модель gpt://<folder_id>/yandexgpt/latest. ИИ может предложить "
        "не более двух типов инструмента с указанием полей-оснований. Сервер проверяет эти основания. ИИ не меняет "
        "обязательные фильтры, не добавляет предложения и не определяет вероятность одобрения. При ошибке, пустом "
        "ответе или таймауте 8 секунд ранжирование идёт по правилам."
    )
    doc.add_paragraph(
        "Карточки сортируются по сумме баллов по убыванию; при равенстве — по ID продукта. "
        "Метка: от 85 — «высокое соответствие», от 70 — «среднее соответствие», ниже 70 — «требует проверки». "
        "В карточке указываются конкретные банк и продукт, диапазоны суммы и ставки, объяснение баллов и предупреждения."
    )
    doc.add_heading("6. Что рейтинг пока не учитывает", level=1)
    doc.add_paragraph(
        "Диапазон ставки показывается, но не входит в формулу. Выручка, юридическая форма и поле "
        "max_debt_to_revenue не участвуют в оценке; просрочка даёт только предупреждение. "
        "Первый транш и отсрочка поставщика формируют подсказки, но не отдельные баллы. "
        "Точный текст нестандартного агрографика сохраняется, однако без классификации ИИ сам по себе рейтинг не меняет."
    )
    doc.add_paragraph(
        "Условия 12 предложений внесены в seed-каталог и имеют статус unverified. Воркер проверяет доступность "
        "страниц банков, но не обновляет и не подтверждает ставки, лимиты и критерии. Баллы — соответствие "
        "текущей модели, не вероятность одобрения и не доказательство, что это самый дешёвый кредит на рынке."
    )
    paragraph(doc, "Источник реализации: ", "app/matching.py, app/main.py, app/ai.py, app/seed.py, app/worker.py.")
    path = OUT / "ranking_akimai.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for generated in (build_questionnaire(), build_ranking()):
        print(generated)
