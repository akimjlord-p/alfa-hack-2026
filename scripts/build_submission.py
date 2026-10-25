"""Generate machine-readable API files and a compact presentation for submission."""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission"
sys.path.insert(0, str(ROOT))


def build_api_specs() -> None:
    from app.main import app
    schema = app.openapi()
    (OUT / "openapi.json").write_text(json.dumps(schema, ensure_ascii=False, indent=2), encoding="utf-8")
    try:
        import yaml
        (OUT / "openapi.yaml").write_text(yaml.safe_dump(schema, allow_unicode=True, sort_keys=False), encoding="utf-8")
    except ImportError:
        (OUT / "openapi.yaml").write_text("# Install PyYAML to regenerate this file; openapi.json is canonical.\n", encoding="utf-8")


def build_presentation() -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    green = RGBColor(13, 96, 69)
    dark = RGBColor(21, 58, 45)
    muted = RGBColor(91, 112, 100)

    def slide(title: str, subtitle: str, bullets: list[str]) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid(); bg.fore_color.rgb = RGBColor(244, 247, 244)
        box = s.shapes.add_textbox(Inches(.7), Inches(.55), Inches(12), Inches(.8))
        tf = box.text_frame; tf.text = title
        tf.paragraphs[0].font.size = Pt(30); tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.color.rgb = green
        sub = s.shapes.add_textbox(Inches(.72), Inches(1.42), Inches(11.8), Inches(.5))
        sub.text_frame.text = subtitle; sub.text_frame.paragraphs[0].font.size = Pt(14); sub.text_frame.paragraphs[0].font.color.rgb = muted
        body = s.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(11.2), Inches(4.5))
        tf = body.text_frame; tf.clear()
        for i, item in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item; p.level = 0; p.font.size = Pt(21); p.font.color.rgb = dark; p.space_after = Pt(16)
        footer = s.shapes.add_textbox(Inches(.72), Inches(7.05), Inches(11.8), Inches(.2))
        footer.text_frame.text = "akimai · демо-прототип · синтетические заявки"
        footer.text_frame.paragraphs[0].font.size = Pt(9); footer.text_frame.paragraphs[0].font.color.rgb = muted

    slide("akimai", "Кредитный кабинет предпринимателя", [
        "Адаптивная анкета → лента продуктов → демо-заявка → обработка и решение.",
        "Цель прототипа — объяснимо сопоставить задачу бизнеса с конкретными кредитными предложениями.",
    ])
    slide("Проблема и решение", "Предприниматель часто не знает, какой инструмент ему нужен", [
        "Система спрашивает про отрасль, денежный цикл, этапность расходов, обеспечение и способность платить.",
        "Вопросы ветвятся: строительство не получает аграрные подсказки, кассовый разрыв ведёт к коротким инструментам.",
        "ИИ извлекает факты из свободного текста, но правила сохраняют контроль над маршрутом и допуском.",
    ])
    slide("Сквозной путь", "От задачи до проверяемого демо-решения", [
        "1 · анкета и нормализованный профиль",
        "2 · фильтры каталога и объяснимый рейтинг",
        "3 · заявка со статусом submitted",
        "4 · worker → under_review или needs_info",
        "5 · эксперт → approved или rejected с событием и комментарием",
    ])
    slide("Архитектура и данные", "FastAPI + PostgreSQL + worker + Yandex AI Studio", [
        "PostgreSQL хранит сессии, сообщения, продукты, заявки и полный таймлайн событий.",
        "Worker приоритизирует пользовательскую первичную проверку над медленной проверкой страниц источников.",
        "OpenAPI и Postman-коллекция лежат в комплекте сдачи.",
    ])
    slide("Как выбирается продукт", "Жёсткие фильтры → баллы → сортировка", [
        "Отсекаются несовпадения цели, суммы и обеспечения.",
        "Баллы: инструмент 22–35, денежный цикл 0–20, каталог 16–20, платёж 0–15, скорость 0–5, ИИ 0–5.",
        "Это соответствие модели, а не вероятность одобрения; предложения помечены unverified.",
    ])
    slide("Результат и границы", "Демо можно повторить без устных пояснений", [
        "Есть роли предпринимателя и эксперта, две заранее подготовленные синтетические заявки и сценарии уточнения.",
        "Есть локальные тесты, браузерный smoke-тест и отдельные схемы архитектуры, состояний и данных.",
        "Не имитируем реальный банк, БКИ, ЭЦП или юридически значимое решение — это явно обозначено в документации.",
    ])
    prs.save(OUT / "presentation.pptx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_api_specs()
    build_presentation()
    print("submission generated")
